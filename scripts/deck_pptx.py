"""deck.json → PPT。版面與配色全部來自 templates/<name>/theme.json（影片 src/deck/DeckVideo.tsx 讀同一份）。

    python scripts/deck_pptx.py <project> [輸出路徑]

頁型：cover／bullets／boxes／steps／screen，欄位見 projects/README.md。
"""
import json
import pathlib
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = pathlib.Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))
from deck import load, template_dir  # noqa: E402

deck, theme, name = load(sys.argv[1])
out = pathlib.Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "projects" / name / "out" / f"{name}.pptx"
TPL = template_dir(deck)
SHOTS = ROOT / "projects" / name / "shots"
L = theme["layout"]
FONT = theme["font"]


def rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


C = {k: rgb(v) for k, v in theme["colors"].items()}
TOTAL = len(deck["slides"])

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(theme["slide"]["w"]), Inches(theme["slide"]["h"])
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, runs, size=18, color=None, bold=False, align=PP_ALIGN.LEFT, spacing=None, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = [[(runs, {})]] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for s, o in para:
            r = p.add_run()
            r.text = s
            f = r.font
            f.name = FONT
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = italic
            f.color.rgb = o.get("color", color or C["body"])
            if spacing:
                r._r.get_or_add_rPr().set("spc", str(spacing))
    return tb


def shape(slide, kind, x, y, w, h, fill=None, line=None, width=1.0):
    s = slide.shapes.add_shape(kind, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(width)
    s.shadow.inherit = False
    return s


def dot(slide, x, y, d, color):
    return shape(slide, MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d), fill=color)


def page(n, bg, kicker=None, title=None):
    s = prs.slides.add_slide(BLANK)
    s.shapes.add_picture(str(TPL / theme["backgrounds"][bg]), 0, 0, prs.slide_width, prs.slide_height)
    p = L["page"]
    if kicker is not None:
        k = p["kicker"]; text(s, k["x"], k["y"], k["w"], 0.35, kicker, size=k["size"], color=C["accent"], spacing=k["spacing"])
    if title is not None:
        t = p["title"]; text(s, t["x"], t["y"], t["w"], 0.75, title, size=t["size"], color=C["white"], bold=True)
    pn = p["pageNo"]
    return s, (lambda: text(s, pn["x"], pn["y"], pn["w"], 0.3, f"{n} / {TOTAL}", size=pn["size"], color=C["body"], align=PP_ALIGN.RIGHT))


def tone(t):
    return C.get(t or "accent", C["accent"])


def cover(n, sl):
    s, pageno = page(n, "cover")
    c = L["cover"]
    if sl.get("year"):
        y = c["year"]; text(s, y["x"], y["y"], y["w"], 2.4, sl["year"], size=y["size"], color=C["digits"], bold=True, italic=True)
    o = c["org"]; text(s, o["x"], o["y"], o["w"], 0.4, sl.get("org", ""), size=o["size"], color=C["body"])
    t = c["title"]; text(s, t["x"], t["y"], t["w"], 0.9, sl["title"], size=t["size"], color=C["coral"], bold=True)
    r = c["rule"]; shape(s, MSO_SHAPE.RECTANGLE, Inches(r["x"]), Inches(r["y"]), Inches(r["w"]), Inches(r["h"]), fill=C["coral"])
    st = c["subtitle"]
    text(s, st["x"], st["y"], st["w"], 0.9, [[(x["text"], {"color": C["white"]} if x.get("strong") else {})] for x in sl.get("subtitle", [])], size=st["size"])
    pageno()


def bullets(n, sl):
    s, pageno = page(n, "content", sl.get("kicker"), sl.get("title"))
    pageno()
    b = L["bullets"]
    for i, t in enumerate(sl["items"]):
        y = b["y"] + i * b["gap"]
        dot(s, b["x"], y + b["dotDy"], b["dot"], tone(sl.get("tone")))
        text(s, b["x"] + b["textDx"], y, b["w"] - b["textDx"], b["gap"], t, size=b["size"])
    if sl.get("note"):
        text(s, b["x"], b["note"]["y"], b["w"], 0.4, sl["note"], size=b["note"]["size"], color=C["label"])


def boxes(n, sl):
    s, pageno = page(n, "content", sl.get("kicker"), sl.get("title"))
    pageno()
    b = L["boxes"]
    for i, bx in enumerate(sl["boxes"]):
        x = b["x"] + i * b["step"]
        shape(s, MSO_SHAPE.RECTANGLE, Inches(x), Inches(b["y"]), Inches(b["w"]), Inches(b["h"]), fill=C["box"], line=C["boxLine"])
        text(s, x + b["padX"], b["headY"], b["w"] - 2 * b["padX"], 0.5, bx["head"], size=b["headSize"], bold=True, color=C["white"])
        text(s, x + b["padX"], b["linesY"], b["w"] - 2 * b["padX"], 1.6, [[(ln, {})] for ln in bx["lines"]], size=b["linesSize"])
        if i < len(sl["boxes"]) - 1:
            text(s, x + b["w"], b["arrowY"], b["arrowW"], 0.5, "→", size=b["arrowSize"], color=C["accent"], align=PP_ALIGN.CENTER)


def steps(n, sl):
    s, pageno = page(n, "content", sl.get("kicker"), sl.get("title"))
    pageno()
    st = L["steps"]
    for (key, label), lx in zip((sl.get("legend") or {}).items(), st["legendX"]):
        dot(s, lx, st["legendDotY"], st["legendDot"], tone(key))
        text(s, lx + st["legendTextDx"], st["legendY"], 1.3, 0.35, label, size=st["legendSize"])
    for i, row in enumerate(sl["steps"]):
        y = st["y"] + i * st["gap"]
        human = row.get("human", False)
        c = C["coral"] if human else C["accent"]
        o = dot(s, st["circleX"], y + st["circleDy"], st["circle"], c)
        o.text_frame.text = str(i + 1)
        p = o.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.runs[0]; r.font.size = Pt(st["numSize"]); r.font.bold = True; r.font.color.rgb = C["white"]; r.font.name = FONT
        text(s, st["whoX"], y + st["whoDy"], st["textX"] - st["whoX"], 0.4, row["who"], size=st["whoSize"], color=c, bold=human)
        text(s, st["textX"], y + st["textDy"], st["descX"] - st["textX"], 0.45, row["text"], size=st["textSize"], color=C["white"], bold=human)
        text(s, st["descX"], y + st["descDy"], st["descW"], 0.5, row.get("desc", ""), size=st["descSize"])
        ln = s.shapes.add_connector(1, Inches(st["circleX"]), Inches(y + st["ruleDy"]), Inches(st["ruleX2"]), Inches(y + st["ruleDy"]))
        ln.line.color.rgb = C["rule"]; ln.line.width = Pt(0.75); ln.shadow.inherit = False


def screen(n, sl):
    s, pageno = page(n, "content", sl.get("kicker"), sl.get("title"))
    pageno()
    sc = L["screen"]
    marks = json.loads((SHOTS / "shots.json").read_text(encoding="utf-8"))[sl["shot"]]
    img_w = sc["imgW"]
    img_h = img_w * sc["shotH"] / sc["shotW"]
    k = Inches(img_w) / sc["shotW"]
    x0, y0 = Inches(sc["imgX"]), Inches(sc["imgY"])
    s.shapes.add_picture(str(SHOTS / f"{sl['shot']}.png"), x0, y0, width=Inches(img_w))
    shape(s, MSO_SHAPE.RECTANGLE, x0, y0, Inches(img_w), Inches(img_h), line=C["boxLine"], width=1)
    tones = sl.get("tones", [])
    t_at = lambda i: tone(tones[i] if i < len(tones) else "accent")
    for i, r in enumerate(marks):
        pad = 4 * k
        h = min(r[1] + r[3], sc["shotH"] - 6) - r[1]   # 框不要超出截圖
        shape(s, MSO_SHAPE.RECTANGLE, x0 + r[0] * k - pad, y0 + r[1] * k - pad, r[2] * k + 2 * pad, h * k + 2 * pad, line=t_at(i), width=2.25)
    bx = sc["imgX"] + img_w + sc["gap"]
    for i, t in enumerate(sl["items"]):
        y = sc["pointY"] + i * sc["pointGap"]
        dot(s, bx, y + sc["dotDy"], sc["dot"], t_at(i))
        text(s, bx + sc["textDx"], y, sc["rightEdge"] - bx - sc["textDx"], 1.0, t, size=sc["pointSize"])


KINDS = {"cover": cover, "bullets": bullets, "boxes": boxes, "steps": steps, "screen": screen}
for n, sl in enumerate(deck["slides"], 1):
    KINDS[sl["type"]](n, sl)
out.parent.mkdir(parents=True, exist_ok=True)
prs.save(out)
print("wrote", out, f"({TOTAL} slides)")
